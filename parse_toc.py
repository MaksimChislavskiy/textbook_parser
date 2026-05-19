import os
import re
import openpyxl
from docx import Document

INPUT_FOLDER = 'input_data'
OUTPUT_FOLDER = 'output'
TARGET_FILE = 'tekstovye_zadachi_po_matematike.docx'

EXCEPTIONS = {
    'Оглавление': 0,
    'Ответы': 0,
    'Приложение': 0,
}

SECTION_RE = re.compile(
    r'^(?P<номер>\d+)\.\s+(?P<название>.+?)\s*(?:[\.\s]{3,})?\s*(?P<страница>\d+)$'
)
SUBSECTION_RE = re.compile(
    r'^(?P<номер>\d+\.\d+)\.?\s+(?P<название>.+?)\s*(?:[\.\s]{3,})?\s*(?P<страница>\d+)$'
)

if TARGET_FILE:
    doc_path = os.path.join(INPUT_FOLDER, TARGET_FILE)
    if not os.path.exists(doc_path):
        raise FileNotFoundError(f'Файл {doc_path} не найден')
    docx_files = [TARGET_FILE]
else:
    docx_files = [f for f in os.listdir(INPUT_FOLDER) if f.endswith('.docx')]
    if not docx_files:
        raise FileNotFoundError(f'Нет .docx файлов в папке {INPUT_FOLDER}/')

os.makedirs(OUTPUT_FOLDER, exist_ok=True)

for filename in docx_files:
    doc = Document(os.path.join(INPUT_FOLDER, filename))
    all_lines = []
    for paragraph in doc.paragraphs:
        line = paragraph.text.strip()
        if line:
            all_lines.append(line)
    data = []
    buffer = ""
    next_id = 1
    current_section_id = 0
    in_toc = False
    for line in all_lines:
        if line == 'Оглавление':
            in_toc = True
            continue
        if not in_toc:
            continue
        sm = SECTION_RE.match(line)
        ssm = SUBSECTION_RE.match(line)
        if re.match(r'^\d+$', line.strip()):
            continue
        if not sm and not ssm and not buffer:
            in_toc = False
            continue
        if sm:
            full_line = (buffer + " " + line).strip() if buffer else line
            buffer = ""
            sm = SECTION_RE.match(full_line)
            if sm:
                name = sm.group('название').strip().rstrip('.')
                if name in EXCEPTIONS:
                    data.append((next_id, name, EXCEPTIONS[name]))
                else:
                    data.append((next_id, name, 0))
                current_section_id = next_id
                next_id += 1
        elif ssm:
            full_line = (buffer + " " + line).strip() if buffer else line
            buffer = ""
            ssm = SUBSECTION_RE.match(full_line)
            sm = SECTION_RE.match(full_line)
            if ssm:
                name = ssm.group('название').strip().rstrip('.')
                if name in EXCEPTIONS:
                    data.append((next_id, name, EXCEPTIONS[name]))
                else:
                    data.append((next_id, name, current_section_id))
                next_id += 1
            elif sm:
                name = sm.group('название').strip().rstrip('.')
                if name in EXCEPTIONS:
                    data.append((next_id, name, EXCEPTIONS[name]))
                else:
                    data.append((next_id, name, 0))
                current_section_id = next_id
                next_id += 1
        else:
            buffer = (buffer + " " + line).strip() if buffer else line

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "table_of_contents"
    ws.append(["id", "название", "parent"])

    for row in data:
        ws.append(row)

    output_name = f"результат_{filename.replace('.docx', '.xlsx')}"
    output_path = os.path.join(OUTPUT_FOLDER, output_name)

    try:
        wb.save(output_path)
        print(f"Готово: {filename} → {output_name} ({len(data)} записей)")
    except PermissionError:
        print(f"Ошибка: закройте файл {output_path} перед повторным запуском")