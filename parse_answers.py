import os
import re
import openpyxl
from docx import Document

INPUT_FOLDER = 'input_data'
OUTPUT_FOLDER = 'output'
TARGET_FILE = 'tekstovye_zadachi_po_matematike.docx'


def clean_answer(text):
    text = re.sub(r'[;\.]\s*$', '', text)
    text = re.sub(r'\s+', ' ', text)
    return text.strip()


def parse_answers_final(text):
    """Финальная рабочая версия парсера"""
    results = []
    text = re.sub(r'^Ответы и советы\.?\s*', '', text)
    
    # Специальная обработка задачи 6 (без пробела после точки)
    task6_match = re.search(r'6\.(\d+\s+[^\.]+\.)', text)
    if task6_match:
        results.append(('6)', clean_answer(task6_match.group(1))))
        text = re.sub(r'6\.\d+\s+[^\.]+\.', '', text)
    
    # Специальная обработка задачи 8.3
    task8_match = re.search(r'8\.\s+3\)\s+(\d+)\.', text)
    if task8_match:
        results.append(('8.3)', task8_match.group(1)))
        text = re.sub(r'8\.\s+3\)\s+\d+\.', '', text)
    
    # Специальная обработка задачи 11.1
    task11_match = re.search(r'11\.\s+1\)\s+([^;\.]+[;\.]\s*[^\.]+)', text)
    if task11_match:
        results.append(('11.1)', clean_answer(task11_match.group(1))))
        text = re.sub(r'11\.\s+1\)\s+[^;\.]+[;\.]\s*[^\.]+\.', '', text)
    
    # Специальная обработка задачи 13 (без пробела после точки)
    task13_match = re.search(r'13\.(\d+\s+[р\.]+)', text)
    if task13_match:
        results.append(('13)', clean_answer(task13_match.group(1))))
        text = re.sub(r'13\.\d+\s+[р\.]+\.', '', text)
    
    # Основной парсинг
    pos = 0
    while pos < len(text):
        match = re.search(r'(\d+)\.\s+', text[pos:])
        if not match:
            break
        
        task_num = match.group(1)
        task_content_start = pos + match.end()
        
        next_match = re.search(r'\d+\.\s+', text[task_content_start:])
        if next_match:
            task_end = task_content_start + next_match.start()
        else:
            task_end = len(text)
        
        task_content = text[task_content_start:task_end].strip()
        
        letter_pattern = r'([а-я]|д)\)\s*([^;]+?)(?=;\s*(?:[а-я]|д)\)|\.\s*$|\.\s*\d+\.)'
        number_pattern = r'(\d+)\)\s*([^;]+?)(?=;\s*(?:\d+\)|[а-я]\))|\.\s*$|\.\s*\d+\.)'
        
        sub_matches = []
        for sub_match in re.finditer(letter_pattern, task_content):
            sub_letter = sub_match.group(1)
            answer = sub_match.group(2).strip()
            answer = clean_answer(answer)
            if answer:
                sub_matches.append((sub_letter, answer))        
        if not sub_matches:
            for sub_match in re.finditer(number_pattern, task_content):
                sub_num = sub_match.group(1)
                answer = sub_match.group(2).strip()
                answer = clean_answer(answer)
                if answer:
                    sub_matches.append((sub_num, answer))        
        if sub_matches:
            for sub, answer in sub_matches:
                results.append((f"{task_num}.{sub})", answer))
        else:
            clean_content = clean_answer(task_content)
            if clean_content:
                results.append((f"{task_num})", clean_content))        
        pos = task_end    
    return results


def extract_answers_from_doc(doc_path):
    doc = Document(doc_path)
    all_text = []
    
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            all_text.append(text)
    
    full_text = ' '.join(all_text)
    
    answers_start = full_text.find('Ответы и советы')
    if answers_start == -1:
        print("Блок 'Ответы и советы' не найден")
        return []
    
    answers_end = full_text.find('Оглавление', answers_start)
    if answers_end == -1:
        answers_end = len(full_text)
    
    answers_block = full_text[answers_start:answers_end]
    
    return parse_answers_final(answers_block)


if __name__ == '__main__':
    doc_path = os.path.join(INPUT_FOLDER, TARGET_FILE)
    if not os.path.exists(doc_path):
        raise FileNotFoundError(f'Файл {doc_path} не найден')
    
    os.makedirs(OUTPUT_FOLDER, exist_ok=True)
    answers = extract_answers_from_doc(doc_path)
    
    # Пост-обработка
    answers = [a for a in answers if a[0] not in ['5)', '8)', '32)']]
    
    if '239)' not in [a[0] for a in answers]:
        answers.append(('239)', '15 пчелок'))
    if '240)' not in [a[0] for a in answers]:
        answers.append(('240)', '2376 дахеканов'))
    
    answers = [('110.1)', a[1]) if a[0] == '110.е)' else a for a in answers]
    
    # Удаление дубликатов и сортировка
    seen = set()
    unique_answers = []
    for ans in answers:
        if ans[0] not in seen:
            seen.add(ans[0])
            unique_answers.append(ans)
    unique_answers.sort(key=lambda x: int(re.match(r'(\d+)', x[0]).group(1)))
    
    # Сохранение в Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "answers"
    ws.append(["id_tasks_book", "answer"])
    for ans_id, answer in unique_answers:
        ws.append([ans_id, answer])
    
    output_path = os.path.join(OUTPUT_FOLDER, f"ответы_{TARGET_FILE.replace('.docx', '.xlsx')}")
    wb.save(output_path)
    print(f"Готово! {len(unique_answers)} ответов сохранено в {output_path}")
